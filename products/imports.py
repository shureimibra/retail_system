"""
Stock from Excel: download a template of the shop's items, fill in the counts, upload it.
Reading a file only builds a plan (the preview). Nothing changes until the plan is applied.
"""
import csv
import io
from decimal import Decimal, InvalidOperation

from django.db import transaction
from django.utils import timezone

from .models import Category, Product, ProductType, ProductVariation, StockImport
from . import services

HEADERS = ['SKU', 'Item', 'Type', 'Size', 'Colour', 'Price', 'In stock']
# Other spellings people use for the same columns.
ALIASES = {'sku': 'sku', 'code': 'sku', 'item': 'item', 'product': 'item', 'name': 'item', 'type': 'type',
           'size': 'size', 'colour': 'colour', 'color': 'colour', 'price': 'price', 'in stock': 'stock',
           'stock': 'stock', 'quantity': 'stock', 'qty': 'stock', 'count': 'stock'}
MAX_ROWS = 5000
MAX_BYTES = 2 * 1024 * 1024


def build_template(shop):
    """An .xlsx of every item, size and colour with today's counts, ready to edit."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.worksheet.datavalidation import DataValidation

    book = Workbook()
    sheet = book.active
    sheet.title = 'Stock'
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    products = Product.objects.filter(shop=shop).select_related('product_type').prefetch_related('variations')
    for product in products.order_by('name'):
        type_name = product.product_type.name if product.product_type else ''
        variations = sorted(product.variations.all(), key=lambda v: (v.option1 or '', v.option2 or ''))
        for v in variations or [None]:
            sheet.append([v.sku if v else '', product.name, type_name, (v.option1 or '') if v else '',
                          (v.option2 or '') if v else '', float(product.base_price), v.stock_quantity if v else None])
    for column, width in zip('ABCDEFG', (16, 32, 14, 10, 14, 12, 10)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = 'A2'
    types = [t.name for t in ProductType.objects.filter(shop=shop)]
    if types and len(','.join(types)) < 250:
        check = DataValidation(type='list', formula1='"' + ','.join(types) + '"', allow_blank=True)
        sheet.add_data_validation(check)
        check.add(f'C2:C{MAX_ROWS + 1}')

    notes = book.create_sheet('How to')
    for line in (
        'Change the numbers in "In stock" to what you counted. Count units put aside for customers too.',
        'A new size or colour: add a row with the item name, size, colour and count. Leave SKU blank.',
        'A new item: add a row with its name, type, price and count. Leave SKU blank.',
        'Leave "In stock" empty to keep a count as it is. Don\'t change the SKU column.',
        'Then upload the file in Products, Import stock. You see every change before anything is saved.',
    ):
        notes.append([line])
    notes.column_dimensions['A'].width = 100
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def read_rows(upload):
    """[(row number, {sku, item, type, size, colour, price, stock})] from an .xlsx or .csv upload."""
    name = (upload.name or '').lower()
    if upload.size > MAX_BYTES:
        raise ValueError('That file is too big. Keep it under 2 MB (about 5,000 rows).')
    if name.endswith('.xlsx'):
        from openpyxl import load_workbook
        try:
            book = load_workbook(upload, read_only=True, data_only=True)
        except Exception:
            raise ValueError("That file couldn't be opened. Save it as an Excel workbook (.xlsx) and try again.")
        sheet = book['Stock'] if 'Stock' in book.sheetnames else book.worksheets[0]
        table = [list(row) for row in sheet.iter_rows(values_only=True)]
    elif name.endswith('.csv'):
        raw = upload.read()
        try:
            text = raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = raw.decode('latin-1')
        table = list(csv.reader(io.StringIO(text)))
    else:
        raise ValueError('Upload the Excel file (.xlsx) or a .csv file.')

    if not table:
        raise ValueError('The file is empty.')
    columns = {}
    for i, header in enumerate(table[0]):
        key = ALIASES.get(str(header or '').strip().lower())
        if key and key not in columns:
            columns[key] = i
    missing = [label for key, label in (('item', 'Item'), ('stock', 'In stock')) if key not in columns]
    if missing:
        raise ValueError(f'The first row must have the column names. Missing: {", ".join(missing)}. '
                         'Download the template to see them.')
    rows = []
    for number, values in enumerate(table[1:], start=2):
        row = {key: _text(values[i]) if i < len(values) else '' for key, i in columns.items()}
        if any(row.values()):
            rows.append((number, row))
    if len(rows) > MAX_ROWS:
        raise ValueError(f'That file has {len(rows)} rows. Split it into files of up to {MAX_ROWS}.')
    return rows


def _text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _number(text, whole=False):
    cleaned = text.replace(',', '').replace('KSh', '').replace('KES', '').strip()
    try:
        value = Decimal(cleaned)
    except InvalidOperation:
        return None
    if value < 0 or (whole and value != value.to_integral_value()):
        return None
    return int(value) if whole else value


def plan_import(shop, rows):
    """
    The preview: what saving would change, and the rows that can't be used. Writes nothing.
    Actions: count (an existing size/colour), new_item, new_option, price.
    """
    types = {t.name.lower(): t for t in ProductType.objects.filter(shop=shop)}
    products = {p.name.lower(): p for p in Product.objects.filter(shop=shop)}
    variations = list(ProductVariation.objects.filter(product__shop=shop).select_related('product')
                      .prefetch_related('holds'))
    by_sku = {v.sku.lower(): v for v in variations}
    by_options = {(v.product_id, (v.option1 or '').lower(), (v.option2 or '').lower()): v for v in variations}

    actions, problems, seen, new_items, prices = [], [], set(), {}, {}
    unchanged = 0
    for number, row in rows:
        def problem(text):
            problems.append({'row': number, 'text': text})

        stock = None
        if row.get('stock', ''):
            stock = _number(row['stock'], whole=True)
            if stock is None:
                problem(f'"In stock" must be a whole number, not "{row["stock"]}".')
                continue
        price = None
        if row.get('price', ''):
            price = _number(row['price'])
            if price is None:
                problem(f'Price must be a number, not "{row["price"]}".')
                continue
        size, colour = row.get('size', '')[:50], row.get('colour', '')[:50]
        name = row.get('item', '')[:200]

        variation = by_sku.get(row.get('sku', '').lower()) if row.get('sku') else None
        product = variation.product if variation else products.get(name.lower())
        if variation is None and product is not None:
            variation = by_options.get((product.id, size.lower(), colour.lower()))
        key = ('v', variation.id) if variation else ('n', name.lower(), size.lower(), colour.lower())
        if key in seen:
            problem(f'{name or "This item"} {size} {colour} is listed twice. Keep one row for it.'.replace('  ', ' '))
            continue
        seen.add(key)

        item_key = product.id if product else f'new:{name.lower()}'
        if price is not None:
            if item_key in prices and prices[item_key] != price:
                problem(f'{name} has two prices ({prices[item_key]:,.0f} and {price:,.0f}). Use one price per item.')
                continue
            prices[item_key] = price

        if variation is not None:
            label = f'{variation.product.name} {variation.label}'.strip()
            if stock is None or stock == variation.stock_quantity:
                unchanged += 1
                continue
            held = variation.held_quantity
            if stock < held:
                problem(f'{label}: {held} {"is" if held == 1 else "are"} put aside for customers, '
                        f'so the count can\'t be less than {held}.')
                continue
            actions.append({'kind': 'count', 'variation': variation.id, 'item': label,
                            'old': variation.stock_quantity, 'new': stock})
            continue

        if not name:
            problem('Item name is missing.')
            continue
        if product is None and name.lower() not in new_items:
            product_type = types.get(row.get('type', '').lower())
            if product_type is None:
                choices = ', '.join(t.name for t in types.values())
                problem(f'{name} is a new item: give its Type ({choices}).' if not row.get('type')
                        else f'"{row["type"]}" is not one of your types ({choices}).')
                continue
            if price is None:
                problem(f'{name} is a new item: give its price.')
                continue
            new_items[name.lower()] = True
            actions.append({'kind': 'new_item', 'name': name, 'type': product_type.id, 'price': str(price),
                            'item': name})
        sku = row.get('sku', '')[:100]
        if sku and ProductVariation.objects.filter(sku__iexact=sku).exists():
            problem(f'SKU {sku} belongs to another item. Leave SKU blank for a new size or colour.')
            continue
        actions.append({'kind': 'new_option', 'product': product.id if product else None, 'name': name,
                        'size': size, 'colour': colour, 'sku': sku, 'new': stock or 0,
                        'item': f'{product.name if product else name} {" / ".join(v for v in (size, colour) if v)}'.strip()})

    for item_key, price in prices.items():
        if isinstance(item_key, int):
            product = next(p for p in products.values() if p.id == item_key)
            if price != product.base_price:
                actions.append({'kind': 'price', 'product': product.id, 'item': product.name,
                                'old': str(product.base_price), 'new': str(price)})
    return {'actions': actions, 'problems': problems, 'unchanged': unchanged, 'summary': summarise(actions)}


def summarise(actions):
    """How many of each change, and how many sizes went up (those may message waiting customers)."""
    counts = {kind: sum(1 for a in actions if a['kind'] == kind) for kind in ('count', 'new_option', 'new_item', 'price')}
    counts['up'] = sum(1 for a in actions if a['kind'] == 'count' and a['new'] > a['old'])
    return counts


@transaction.atomic
def apply_import(stock_import, user):
    """Save a previewed import. Counts go through the stock services, so waiting customers are told."""
    stock_import = StockImport.objects.select_for_update().get(pk=stock_import.pk)
    if stock_import.applied_at:
        raise ValueError('This file was already saved.')
    shop = stock_import.shop
    reference = f'Stock import #{stock_import.pk}'
    made, skipped = {}, []
    for action in stock_import.plan['actions']:
        kind = action['kind']
        if kind == 'count':
            variation = ProductVariation.objects.filter(pk=action['variation'], product__shop=shop).first()
            if variation is None:
                skipped.append(f'{action["item"]} was deleted')
                continue
            if action['new'] < variation.held_quantity:
                skipped.append(f'{action["item"]}: more units were put aside since the preview')
                continue
            services.set_stock(variation.pk, action['new'], recorded_by=user, reference=reference)
        elif kind == 'new_item':
            product_type = ProductType.objects.filter(shop=shop, pk=action['type']).first()
            category, _ = Category.objects.get_or_create(shop=shop, name=product_type.name if product_type else 'Other',
                                                         defaults={'product_type': product_type})
            made[action['name'].lower()] = Product.objects.create(
                shop=shop, category=category, product_type=product_type, name=action['name'],
                base_price=Decimal(action['price']))
        elif kind == 'new_option':
            product = (Product.objects.filter(shop=shop, pk=action['product']).first() if action['product']
                       else made.get(action['name'].lower()))
            if product is None:
                skipped.append(f'{action["item"]}: the item was deleted')
                continue
            size, colour = action['size'] or None, action['colour'] or None
            if ProductVariation.objects.filter(product=product, option1__iexact=size or None,
                                               option2__iexact=colour or None).exists():
                skipped.append(f'{action["item"]} was added since the preview')
                continue
            sku = action['sku'] if action['sku'] and not ProductVariation.objects.filter(sku__iexact=action['sku']).exists() \
                else services.suggest_sku(product, size, colour)
            services.create_product_variation(product, size, colour, action['new'], sku, recorded_by=user)
        elif kind == 'price':
            product = Product.objects.filter(shop=shop, pk=action['product']).first()
            if product is not None:
                services.change_price(product, Decimal(action['new']), changed_by=user)
    stock_import.applied_at = timezone.now()
    stock_import.applied_by = user
    stock_import.skipped = skipped
    stock_import.save(update_fields=['applied_at', 'applied_by', 'skipped'])
    return stock_import
