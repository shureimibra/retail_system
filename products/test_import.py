"""Phase 5: stock from Excel. Download the template, upload counts, preview, save."""
from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import Workbook, load_workbook

from accounts.models import Membership, Shop, User
from customers.models import Customer
from intents import services as intents
from intents.models import BlockedPurchaseIntent as Intent
from notifications.models import Message
from .models import Category, Product, ProductType, ProductVariation, StockImport, StockMovement
from .services import receive_stock


def xlsx(rows, headers=('SKU', 'Item', 'Type', 'Size', 'Colour', 'Price', 'In stock')):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Stock'
    sheet.append(list(headers))
    for row in rows:
        sheet.append(list(row))
    buffer = BytesIO()
    book.save(buffer)
    return SimpleUploadedFile('count.xlsx', buffer.getvalue())


class ImportCase(TestCase):
    def setUp(self):
        self.shop = Shop.objects.create(name='Kimathi Street Store')
        self.owner = User.objects.create_user('owner', password='pw12345!')
        Membership.objects.create(user=self.owner, shop=self.shop, role=Membership.OWNER)
        self.client.login(username='owner', password='pw12345!')
        self.shoes = ProductType.objects.get(shop=self.shop, name='Shoes')
        running = Category.objects.create(shop=self.shop, name='Running', product_type=self.shoes)
        self.boost = Product.objects.create(shop=self.shop, category=running, product_type=self.shoes,
                                            name='Adidas Ultraboost', base_price=18700)
        self.grey42 = ProductVariation.objects.create(product=self.boost, option1='42', option2='Grey', sku='AU-42-GRE')
        self.grey44 = ProductVariation.objects.create(product=self.boost, option1='44', option2='Grey', sku='AU-44-GRE',
                                                      stock_quantity=3)

    def upload(self, rows, **kwargs):
        response = self.client.post('/products/import/', {'file': xlsx(rows, **kwargs)})
        return response, StockImport.objects.filter(shop=self.shop).first()


class TemplateTests(ImportCase):
    def test_template_lists_every_size_with_its_count(self):
        response = self.client.get('/products/import/template/')
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(BytesIO(response.content))['Stock']
        rows = [row for row in sheet.iter_rows(values_only=True)]
        self.assertEqual(rows[0], ('SKU', 'Item', 'Type', 'Size', 'Colour', 'Price', 'In stock'))
        self.assertIn(('AU-44-GRE', 'Adidas Ultraboost', 'Shoes', '44', 'Grey', 18700, 3), rows)

    def test_cashiers_cannot_import(self):
        cashier = User.objects.create_user('cash', password='pw12345!')
        Membership.objects.create(user=cashier, shop=self.shop, role=Membership.CASHIER)
        self.client.login(username='cash', password='pw12345!')
        self.assertEqual(self.client.get('/products/import/').status_code, 403)
        self.assertEqual(self.client.get('/products/import/template/').status_code, 403)
        self.assertNotIn('Import stock', self.client.get('/products/').content.decode())


class PreviewTests(ImportCase):
    def test_preview_changes_nothing_until_saved(self):
        response, stock_import = self.upload([
            ('AU-42-GRE', 'Adidas Ultraboost', 'Shoes', '42', 'Grey', 18700, 2),
            ('AU-44-GRE', 'Adidas Ultraboost', 'Shoes', '44', 'Grey', 18700, 3),   # same as now
            ('', 'Adidas Ultraboost', '', '45', 'Black', '', 1),                   # new size/colour
            ('', 'Puma Suede', 'Shoes', '41', 'Red', '6,500', 4),                  # new item
        ])
        self.assertRedirects(response, f'/products/import/{stock_import.pk}/')
        self.grey42.refresh_from_db()
        self.assertEqual(self.grey42.stock_quantity, 0)
        self.assertFalse(Product.objects.filter(name='Puma Suede').exists())
        self.assertEqual(stock_import.plan['summary'], {'count': 1, 'new_option': 2, 'new_item': 1, 'price': 0, 'up': 1})
        self.assertEqual(stock_import.plan['unchanged'], 1)
        page = self.client.get(f'/products/import/{stock_import.pk}/').content.decode()
        self.assertIn('Save these changes', page)
        self.assertIn('Puma Suede 41 / Red', page)

    def test_rows_that_cannot_be_used_are_listed_by_row(self):
        _, stock_import = self.upload([
            ('AU-42-GRE', 'Adidas Ultraboost', 'Shoes', '42', 'Grey', 18700, 'two'),
            ('', 'Mystery', '', '', '', 100, 1),
            ('', 'Mystery 2', 'Shoes', '', '', '', 1),
            ('AU-44-GRE', 'Adidas Ultraboost', '', '44', 'Grey', 18700, 5),
            ('AU-44-GRE', 'Adidas Ultraboost', '', '44', 'Grey', 18700, 6),
        ])
        problems = {p['row']: p['text'] for p in stock_import.plan['problems']}
        self.assertIn('whole number', problems[2])
        self.assertIn('give its Type', problems[3])
        self.assertIn('give its price', problems[4])
        self.assertIn('listed twice', problems[6])
        self.assertEqual(len(stock_import.plan['actions']), 1)  # the good row still counts

    def test_missing_columns_and_wrong_files_are_explained(self):
        response = self.client.post('/products/import/', {'file': xlsx([('x',)], headers=('Name',))}, follow=True)
        self.assertContains(response, 'Missing: In stock')
        response = self.client.post('/products/import/', {'file': SimpleUploadedFile('a.pdf', b'%PDF')}, follow=True)
        self.assertContains(response, 'Upload the Excel file')

    def test_csv_works_too(self):
        csv = SimpleUploadedFile('count.csv', 'Item,Size,Colour,Qty\nAdidas Ultraboost,42,grey,4\n'.encode())
        self.client.post('/products/import/', {'file': csv})
        stock_import = StockImport.objects.get()
        self.assertEqual(stock_import.plan['actions'][0]['new'], 4)

    def test_count_cannot_take_away_units_held_for_customers(self):
        customer = Customer.objects.create(shop=self.shop, first_name='Jane', phone_number='+254711000001')
        intents.create_blocked_purchase_intent(customer, self.boost, self.grey42, Intent.SIZE_MISSING)
        with self.captureOnCommitCallbacks(execute=True):
            receive_stock(self.grey42.pk, 1)  # held for Jane
        _, stock_import = self.upload([('AU-42-GRE', 'Adidas Ultraboost', '', '42', 'Grey', '', 0)])
        self.assertIn('put aside for customers', stock_import.plan['problems'][0]['text'])


class SaveTests(ImportCase):
    def test_saving_applies_counts_and_messages_waiting_customers(self):
        customer = Customer.objects.create(shop=self.shop, first_name='Jane', phone_number='+254711000001')
        intent = intents.create_blocked_purchase_intent(customer, self.boost, self.grey42, Intent.SIZE_MISSING)
        _, stock_import = self.upload([
            ('AU-42-GRE', 'Adidas Ultraboost', 'Shoes', '42', 'Grey', 18700, 2),
            ('AU-44-GRE', 'Adidas Ultraboost', 'Shoes', '44', 'Grey', 18700, 1),
            ('', 'Puma Suede', 'Shoes', '41', 'Red', 6500, 4),
            ('', 'Puma Suede', 'Shoes', '42', 'Red', 6500, ''),
        ])
        with self.captureOnCommitCallbacks(execute=True):
            response = self.client.post(f'/products/import/{stock_import.pk}/', {'action': 'save'}, follow=True)
        self.assertContains(response, 'Stock saved')
        self.grey42.refresh_from_db()
        self.grey44.refresh_from_db()
        self.assertEqual((self.grey42.stock_quantity, self.grey44.stock_quantity), (2, 1))
        puma = Product.objects.get(name='Puma Suede')
        self.assertEqual(sorted((v.option1, v.stock_quantity) for v in puma.variations.all()), [('41', 4), ('42', 0)])
        intent.refresh_from_db()
        self.assertEqual(intent.status, Intent.NOTIFIED)
        self.assertTrue(Message.objects.filter(intent=intent).exists())
        self.assertTrue(StockMovement.objects.filter(reference=f'Stock import #{stock_import.pk}').exists())
        # Saving twice does nothing more.
        response = self.client.post(f'/products/import/{stock_import.pk}/', {'action': 'save'}, follow=True)
        self.assertContains(response, 'already saved')

    def test_a_lower_price_in_the_file_is_a_price_drop(self):
        customer = Customer.objects.create(shop=self.shop, first_name='Ali', phone_number='+254711000002')
        intent = intents.create_blocked_purchase_intent(customer, self.boost, None, Intent.PRICE_TOO_HIGH)
        _, stock_import = self.upload([('AU-44-GRE', 'Adidas Ultraboost', '', '44', 'Grey', '17,500', '')])
        self.assertEqual(stock_import.plan['summary']['price'], 1)
        with self.captureOnCommitCallbacks(execute=True):
            self.client.post(f'/products/import/{stock_import.pk}/', {'action': 'save'})
        self.boost.refresh_from_db()
        intent.refresh_from_db()
        self.assertEqual(self.boost.base_price, 17500)
        self.assertEqual(intent.status, Intent.NOTIFIED)

    def test_cancel_changes_nothing(self):
        _, stock_import = self.upload([('AU-42-GRE', 'Adidas Ultraboost', '', '42', 'Grey', '', 9)])
        self.client.post(f'/products/import/{stock_import.pk}/', {'action': 'cancel'})
        self.grey42.refresh_from_db()
        self.assertEqual(self.grey42.stock_quantity, 0)
        self.assertFalse(StockImport.objects.exists())

    def test_another_shops_import_is_not_reachable(self):
        other = Shop.objects.create(name='Other')
        foreign = StockImport.objects.create(shop=other, file_name='x.xlsx', plan={'actions': []})
        self.assertEqual(self.client.get(f'/products/import/{foreign.pk}/').status_code, 404)
